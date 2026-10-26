"""Native, editable PowerPoint files: executive brand presentation + slide template.

Text boxes are real PowerPoint text set right-to-left in Tajawal (install the
fonts from 04_Typography/ for correct rendering). Logos and screens are images.
"""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt, Emu

import tokens as T
from common import OUT

C = {k: RGBColor.from_string(v["hex"].lstrip("#")) for k, v in T.COLORS.items()}
LOGO = OUT / "02_Logos" / "PNG"
DA = OUT / "10_Digital_Applications"


def _rtl(paragraph, rtl=True):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("rtl", "1" if rtl else "0")


def text(slide, x, y, w, h, runs, size=20, color="ink", bold=False, align="r", rtl=True, anchor="t", font="Tajawal", spacing=1.15):
    """runs: str or list of (text, {size,color,bold}) paragraphs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    paras = runs if isinstance(runs, list) else [(runs, {})]
    for i, (t, o) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"r": PP_ALIGN.RIGHT, "l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER}[o.get("align", align)]
        p.line_spacing = spacing
        _rtl(p, o.get("rtl", rtl))
        r = p.add_run()
        r.text = t
        f = r.font
        f.name = font
        f.size = Pt(o.get("size", size))
        f.bold = o.get("bold", bold)
        f.color.rgb = C[o.get("color", color)] if isinstance(o.get("color", color), str) else o.get("color", color)
        rPr = r._r.get_or_add_rPr()
        for tag in ("a:cs", "a:ea"):
            el = rPr.find(qn(tag))
            if el is None:
                el = rPr.makeelement(qn(tag), {})
                rPr.append(el)
            el.set("typeface", font)
        rPr.set("lang", "ar-SA" if o.get("rtl", rtl) else "en-US")
        if o.get("space_after"):
            p.space_after = Pt(o["space_after"])
    return tb


def rect(slide, x, y, w, h, color, radius=None, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = C[color] if isinstance(color, str) else color
    if line:
        shp.line.color.rgb = C[line]
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    if radius:
        shp.adjustments[0] = radius
    shp.shadow.inherit = False
    return shp


def pic(slide, path, x, y, w=None, h=None):
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)


def new(prs, bg="white"):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = C[bg]
    return s


def chrome(s, n, label):
    text(s, 0.6, 7.02, 6, 0.3, f"ضماناتي · {label}", 10, "stone")
    text(s, 12.2, 7.02, 0.6, 0.3, f"{n:02d}", 10, "ink", True, "l", False)


def heading(s, ar, en, y=0.55):
    text(s, 0.6, y, 12.1, 0.8, ar, 34, "ink", True)
    text(s, 0.6, y + 0.78, 12.1, 0.4, en, 14, "stone", False, "r", False)


def pattern(s, y=4.6, h=2.9, opacity_img="damanati_pattern_white-on-green.svg"):
    pass  # vector pattern is not supported natively; slides use the PNG pattern instead


def presentation():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    patt = OUT / "05_Graphic_Elements" / "patterns" / "damanati_pattern_fade-white-on-green.png"

    # 1 cover
    s = new(prs, "green-600")
    pic(s, patt, 0, 4.2, 13.333)
    pic(s, LOGO / "damanati_bilingual_reversed-on-green_512h.png", 0.6, 0.6, h=0.6)
    text(s, 0.6, 1.9, 12.1, 1.6, [("هوية ضماناتي 2.0", {"size": 54, "bold": True, "color": "white"}),
                                  ("Damanati Brand Identity · Executive presentation", {"size": 18, "color": "white", "rtl": False})])
    # 2 agenda
    s = new(prs)
    heading(s, "المحتوى", "Agenda")
    items = ["التدقيق: ما حافظنا عليه وما صححناه", "جوهر العلامة والقيم", "نظام الشعار", "الألوان والخطوط", "اللغة البصرية",
             "التطبيقات الرقمية والتسويقية", "الحركة", "ما نحتاج اعتماده"]
    for i, t in enumerate(items):
        y = 1.9 + i * 0.6
        text(s, 11.3, y, 1.4, 0.5, f"{i + 1:02d}", 22, "green-600", True, "r", False)
        text(s, 3.5, y, 7.8, 0.5, t, 20, "ink")
    chrome(s, 2, "عرض الهوية")
    # 3 audit
    s = new(prs)
    heading(s, "التدقيق: حافظنا على الهوية، وصححنا التفاصيل", "Audit: identity preserved, details corrected")
    rect(s, 6.9, 1.9, 5.8, 4.6, "paper", 0.06)
    rect(s, 0.6, 1.9, 5.9, 4.6, "green-600", 0.06)
    text(s, 7.2, 2.1, 5.2, 4.2, [("ما بقي كما هو", {"size": 20, "bold": True, "color": "green-700", "space_after": 8}),
                                ("الشعار (ض + D) مستخرج متجهيًا بدون إعادة رسم", {"size": 15}), ("الأخضر #2F7D6B ونسبة 60/30/5/5", {"size": 15}),
                                ("خط تجوال والأوزان السبعة", {"size": 15}), ("النمط، مساحة الحماية، الحد الأدنى 10 ملم / 48px", {"size": 15}),
                                ("الرؤية والرسالة والقيم السبع", {"size": 15})], spacing=1.5)
    text(s, 0.9, 2.1, 5.3, 4.2, [("ما صححناه", {"size": 20, "bold": True, "color": "white", "space_after": 8}),
                                ("ثلاث قيم مختلفة للأخضر → قيمة واحدة", {"size": 15, "color": "white"}), ("guarantees → warranties", {"size": 15, "color": "white"}),
                                ("ثلاثة خطوط لاتينية → تجوال فقط", {"size": 15, "color": "white"}), ("اسم «D-Data» في نموذج الجوال → ضماناتي", {"size": 15, "color": "white"}),
                                ("تعريف مساحة الحماية: x = ⅙ عرض المربع", {"size": 15, "color": "white"})], spacing=1.5)
    chrome(s, 3, "عرض الهوية")
    # 4 essence
    s = new(prs)
    heading(s, "جوهر العلامة", "Brand essence (proposed)")
    for i, (lbl, big, sub) in enumerate([("الغاية", "ما يضيع حق أحد لأن ورقة ضاعت.", "Purpose"),
                                         ("الوعد", "نحفظ ضمانك، ونذكّرك بحقك.", "Promise"),
                                         ("النموذج", "الراعي، بلمسة الحكيم", "Caregiver + Sage")]):
        x = 8.75 - i * 4.1
        rect(s, x, 2.0, 3.95, 3.9, ["green-600", "green-700", "green-800"][i], 0.06)
        text(s, x + 0.3, 2.25, 3.35, 3.4, [(lbl, {"size": 14, "color": "white", "space_after": 10}), (big, {"size": 26, "bold": True, "color": "white", "space_after": 14}),
                                           (sub, {"size": 13, "color": "white", "rtl": False, "align": "r"})], spacing=1.3)
    chrome(s, 4, "عرض الهوية")
    # 5 values
    s = new(prs)
    heading(s, "القيم السبع", "Seven values (approved names)")
    vals = ["الثقة", "الأمان", "الابتكار", "البساطة", "الاعتمادية", "الخصوصية", "التنظيم"]
    col_x = lambda c: 12.7 - (c + 1) * 2.85 - c * 0.23
    for i, v in enumerate(vals):
        x, y = col_x(i % 4), 2.0 + (i // 4) * 2.2
        rect(s, x, y, 2.85, 1.9, "paper", 0.08)
        text(s, x + 0.1, y + 0.55, 2.65, 0.8, v, 26, "ink", True, "c")
    rect(s, col_x(3), 4.2, 2.85, 1.9, "green-600", 0.08)
    text(s, col_x(3) + 0.15, 4.35, 2.55, 1.6, "سبع قيم، وفكرة وحدة: حقك محفوظ.", 18, "white", True, "c", anchor="m")
    chrome(s, 5, "عرض الهوية")
    # 6 logo system
    s = new(prs)
    heading(s, "نظام الشعار", "Logo system: 5 lockups × 5 colourways, extracted as vectors")
    pic(s, LOGO / "damanati_bilingual_primary_2048h.png", 4.4, 2.0, h=0.95)
    pic(s, LOGO / "damanati_symbol_primary_512h.png", 10.6, 3.5, h=1.7)
    pic(s, LOGO / "damanati_arabic_primary_512h.png", 6.2, 3.65, h=1.2)
    pic(s, LOGO / "damanati_english_primary_512h.png", 2.0, 3.65, h=1.2)
    rect(s, 0.6, 5.55, 12.1, 1.15, "green-600", 0.12)
    pic(s, LOGO / "damanati_bilingual_reversed-on-green_512h.png", 4.9, 5.83, h=0.6)
    chrome(s, 6, "عرض الهوية")
    # 7 colour
    s = new(prs)
    heading(s, "الألوان", "Colour: one green leads")
    x = 12.7
    for k, w in [("green-600", 6.6), ("white", 3.3), ("stone", 1.1), ("ink", 1.1)]:
        x -= w
        rect(s, x, 1.9, w - 0.08, 3.6, k, 0.04, line="sand" if k == "white" else None)
        v = T.COLORS[k]
        fg = "white" if T.luminance(v["hex"]) < .4 else "ink"
        text(s, x + 0.1, 4.6, w - 0.3, 0.8, [(v["hex"], {"size": 13 if w > 2 else 10, "bold": True, "color": fg, "rtl": False, "align": "l"})])
    text(s, 0.6, 5.75, 12.1, 0.9, "60% أخضر · 30% أبيض · 5% رمادي · 5% أسود. ألوان الحالة (ساري / ينتهي قريبًا / منتهي) للواجهات فقط.", 16, "stone")
    chrome(s, 7, "عرض الهوية")
    # 8 type
    s = new(prs)
    heading(s, "الخط: تجوال", "Typography: Tajawal (SIL OFL 1.1)")
    text(s, 6.6, 1.9, 6.1, 2.4, "ض Aa", 120, "green-600", True)
    text(s, 0.6, 2.1, 5.8, 4.4, [("كل ضماناتك في مكان واحد", {"size": 30, "bold": True, "space_after": 6}),
                                ("All your warranties in one place", {"size": 16, "color": "stone", "rtl": False, "align": "r", "space_after": 18}),
                                ("خط واحد للعربية واللاتينية، سبعة أوزان من ExtraLight إلى Black، وسلم طباعي من 12 إلى 56 بكسل.", {"size": 16})], spacing=1.4)
    chrome(s, 8, "عرض الهوية")
    # 9 visual families
    s = new(prs)
    heading(s, "اللغة البصرية: ثلاث عائلات", "Visual language: three composition families")
    for i, (img, t) in enumerate([("13_Mockups/campaign_key_visual_16x9.png", "جواز المنتج"),
                                  ("12_Social_Media/instagram_post_03_feature.png", "قوس دورة الحياة"),
                                  ("12_Social_Media/x_post_1600x900.png", "النمط المختوم")]):
        x = 8.75 - i * 4.1
        p = OUT / img
        if "instagram" in img:
            pic(s, p, x + 0.7, 1.95, h=2.5)
        else:
            pic(s, p, x, 2.15, w=3.95)
        text(s, x, 4.75, 3.95, 0.5, t, 18, "ink", True, "c")
    chrome(s, 9, "عرض الهوية")
    # 10 digital
    s = new(prs, "paper")
    heading(s, "داخل المنتج", "Digital applications")
    for i, n in enumerate(["app_01_launch", "app_04_dashboard_light", "app_06_warranty_detail", "app_05_dashboard_dark"]):
        pic(s, DA / f"{n}.png", 10.3 - i * 3.0, 1.75, h=5.1)
    chrome(s, 10, "عرض الهوية")
    # 11 marketing
    s = new(prs)
    heading(s, "التسويق: «لكل منتج جواز ضمان»", "Launch campaign key visual")
    pic(s, OUT / "13_Mockups/campaign_key_visual_16x9.png", 2.6, 1.85, w=8.1)
    chrome(s, 11, "عرض الهوية")
    # 12 corporate
    s = new(prs, "paper")
    heading(s, "الهوية المؤسسية", "Corporate identity")
    pic(s, OUT / "11_Corporate_Identity/letterhead_a4.png", 9.6, 1.8, h=4.9)
    pic(s, OUT / "11_Corporate_Identity/business_card_front.png", 5.2, 1.9, w=3.6)
    pic(s, OUT / "11_Corporate_Identity/business_card_back.png", 5.2, 4.4, w=3.6)
    pic(s, OUT / "11_Corporate_Identity/staff_id_card.png", 1.4, 1.95, h=4.4)
    chrome(s, 12, "عرض الهوية")
    # 13 motion
    s = new(prs, "green-600")
    text(s, 0.6, 0.55, 12.1, 0.8, "الحركة: كشف «الختم»", 34, "white", True)
    text(s, 0.6, 1.3, 12.1, 0.4, "Logo reveal, rendered: 09_Motion_Assets/damanati_logo_reveal_seal_1920x1080.mp4", 14, "white", False, "r", False)
    kf = OUT / "09_Motion_Assets/keyframes"
    for i, t in enumerate(["0.2", "0.7", "1.1", "1.6", "2.2", "3.5"]):
        x = 8.75 - (i % 3) * 4.1
        y = 2.0 + (i // 3) * 2.45
        pic(s, kf / f"seal_t{t}s.png", x, y, w=3.95)
    # 14 approvals
    s = new(prs)
    heading(s, "ما نحتاج اعتماده", "Decisions needed from the brand owner")
    asks = ["اعتماد الأخضر #2F7D6B قيمة وحيدة وتحديث كل الملفات", "اعتماد المصطلح الإنجليزي warranty", "اعتماد الصيغة العمودية والأحادية للشعار",
            "اعتماد ألوان الحالة والثيم الداكن", "اعتماد الغاية والوعد والنبرة", "اختبار طباعة CMYK ومطابقة Pantone", "تحقق ميداني من شخصيات الجمهور"]
    for i, a in enumerate(asks):
        y = 1.9 + i * 0.66
        rect(s, 12.3, y + 0.1, 0.36, 0.36, "green-600", 0.25)
        text(s, 1.0, y, 11.1, 0.55, a, 19, "ink")
    chrome(s, 14, "عرض الهوية")
    # 15 close
    s = new(prs, "green-600")
    pic(s, patt, 0, 4.2, 13.333)
    pic(s, LOGO / "damanati_vertical_reversed-on-green_512h.png", 5.45, 1.2, h=2.6)
    text(s, 0.6, 4.0, 12.1, 0.6, "نحفظ ضمانك، ونذكّرك بحقك.", 22, "white", False, "c")
    out = OUT / "01_Brand_Book" / "Damanati_Brand_Presentation.pptx"
    prs.save(out)
    return out


def template():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    patt = OUT / "05_Graphic_Elements" / "patterns" / "damanati_pattern_fade-white-on-green.png"
    s = new(prs, "green-600")
    pic(s, patt, 0, 4.2, 13.333)
    pic(s, LOGO / "damanati_bilingual_reversed-on-green_512h.png", 0.6, 0.6, h=0.6)
    text(s, 0.6, 2.0, 12.1, 1.6, [("عنوان العرض", {"size": 54, "bold": True, "color": "white"}), ("العنوان الفرعي · التاريخ", {"size": 20, "color": "white"})])
    s = new(prs, "paper")
    text(s, 0.6, 2.0, 12.1, 1.4, "01", 110, "green-600", True)
    text(s, 0.6, 3.95, 12.1, 0.9, "عنوان القسم", 44, "ink", True)
    text(s, 0.6, 4.8, 12.1, 0.5, "Section title", 18, "stone", False, "r", False)
    s = new(prs)
    text(s, 0.6, 0.45, 9, 0.4, "01 · اسم القسم", 14, "green-600")
    pic(s, LOGO / "damanati_symbol_primary_512h.png", 0.6, 0.4, h=0.55)
    text(s, 0.6, 1.2, 12.1, 1.0, "عنوان الشريحة في سطر واحد", 36, "ink", True)
    for i in range(3):
        x = 8.75 - i * 4.1
        rect(s, x, 2.6, 3.95, 0.06, "green-600")
        text(s, x, 2.8, 3.95, 3.5, [(f"0{i + 1}", {"size": 32, "bold": True, "color": "green-600"}), ("عنوان النقطة", {"size": 20, "bold": True}),
                                    ("نص داعم قصير لا يتجاوز ثلاثة أسطر.", {"size": 15, "color": "stone"})], spacing=1.3)
    chrome(s, 3, "قالب العرض")
    s = new(prs, "green-600")
    pic(s, LOGO / "damanati_vertical_reversed-on-green_512h.png", 5.45, 1.6, h=2.6)
    text(s, 0.6, 4.5, 12.1, 0.6, "شكرًا", 28, "white", True, "c")
    out = OUT / "14_Templates" / "Damanati_Presentation_Template.pptx"
    prs.save(out)
    return out


def build():
    # PNG pattern for slides (PowerPoint cannot use the CSS-masked SVG)
    import json
    import subprocess
    from common import BUILD, SRC
    BUILD.mkdir(exist_ok=True)
    svg = OUT / "05_Graphic_Elements" / "patterns" / "damanati_pattern_fade-white-on-green.svg"
    page = BUILD / "patt.html"
    page.write_text(f'<html><body style="margin:0"><img src="{svg.as_uri()}" style="width:2560px;height:560px;object-fit:cover;object-position:top;display:block"></body></html>')
    job = BUILD / "jobs_patt.json"
    job.write_text(json.dumps([{"src": str(page), "out": str(svg.with_suffix(".png")), "width": 2560, "height": 560}]))
    subprocess.run(["node", str(SRC / "render.mjs"), str(job)], check=True, stdout=subprocess.DEVNULL)
    print(presentation().name, template().name)


if __name__ == "__main__":
    build()
